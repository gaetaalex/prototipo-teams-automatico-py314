from __future__ import annotations

import html
import base64
import os
import re
import secrets
import shutil
import unicodedata
import webbrowser
from datetime import date, datetime, time, timedelta
from pathlib import Path

import msal
import requests
from dotenv import load_dotenv
from flask import Flask, flash, redirect, render_template_string, request, send_file, session, url_for
from openpyxl import load_workbook
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "dados"
DATA_DIR.mkdir(exist_ok=True)
load_dotenv(BASE_DIR / ".env")

PORT = int(os.getenv("PORT", "5056"))
TENANT_ID = os.getenv("TENANT_ID", "").strip()
CLIENT_ID = os.getenv("CLIENT_ID", "").strip()
CLIENT_SECRET = os.getenv("CLIENT_SECRET", "").strip()
REDIRECT_URI = os.getenv("REDIRECT_URI", f"http://localhost:{PORT}/auth/callback").strip()
SCOPES = ["User.Read", "User.Read.All", "Chat.Create", "ChatMessage.Send", "Files.ReadWrite"]
GRAPH = "https://graph.microsoft.com/v1.0"
ONEDRIVE_SHARE_URL = os.getenv("ONEDRIVE_SHARE_URL", "").strip()

app = Flask(__name__)
app.secret_key = os.getenv("APP_SECRET", secrets.token_hex(32))
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024


@app.before_request
def padronizar_host_local():
    # OAuth depende do cookie de sessão. 127.0.0.1 e localhost são hosts
    # diferentes para o navegador; usar ambos provoca "state mismatch".
    if request.host.split(":", 1)[0] == "127.0.0.1":
        destino = f"http://localhost:{PORT}{request.full_path}"
        if destino.endswith("?"):
            destino = destino[:-1]
        return redirect(destino, code=302)

# Protótipo local: cache somente em memória. Em produção, usar cache servidor
# criptografado (Redis/SQL), nunca tokens em cookies ou no Excel.
TOKEN_CACHES: dict[str, msal.SerializableTokenCache] = {}

ALIASES = {
    "matricula": {"matricula", "matricula usuario", "id matricula"},
    "nome": {"nome", "usuario", "nome usuario"},
    "mensagem": {"texto de contato", "mensagem", "texto contato"},
    "mensagem2": {"texto segundo contato", "mensagem segundo contato", "segunda mensagem"},
    "status": {"status envio", "status de envio", "status"},
    "numero": {"numero de contato", "numero contato", "n de contato", "contato numero"},
    "ultima_data": {"ultima data de contato", "ultima data contato", "data ultimo contato"},
}


def configurado() -> bool:
    invalidos = {"", "COLE_AQUI", "PREENCHER_COM_A_TI"}
    return TENANT_ID not in invalidos and CLIENT_ID not in invalidos and CLIENT_SECRET not in invalidos


def normalizar(valor) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-zA-Z0-9]+", " ", texto).strip().lower()


def vazio(valor) -> bool:
    return valor is None or str(valor).strip() == ""


def numero_inteiro(valor) -> int:
    try:
        return 0 if vazio(valor) else int(float(valor))
    except (TypeError, ValueError):
        return 0


def converter_data(valor):
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, time.min)
    if vazio(valor):
        return None
    for formato in ("%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(valor).strip(), formato)
        except ValueError:
            pass
    return None


def localizar_cabecalho(ws):
    melhor = None
    for linha in range(1, min(ws.max_row, 50) + 1):
        colunas = {}
        for coluna in range(1, ws.max_column + 1):
            titulo = normalizar(ws.cell(linha, coluna).value)
            for chave, nomes in ALIASES.items():
                if titulo in nomes and chave not in colunas:
                    colunas[chave] = coluna
        pontos = sum(k in colunas for k in ("matricula", "nome", "mensagem", "status", "numero", "ultima_data"))
        if melhor is None or pontos > melhor[0]:
            melhor = (pontos, linha, colunas)
    if not melhor or melhor[0] < 4:
        raise ValueError("Não encontrei uma linha de cabeçalho compatível.")
    return melhor[1], melhor[2]


def abrir_planilha(data_only=False):
    caminho = session.get("planilha")
    if not caminho or not Path(caminho).exists():
        raise FileNotFoundError("Carregue a planilha primeiro.")
    wb = load_workbook(caminho, data_only=data_only)
    aba = session.get("aba")
    if aba and aba in wb.sheetnames:
        ws = wb[aba]
        cabecalho, colunas = localizar_cabecalho(ws)
        return Path(caminho), wb, ws, cabecalho, colunas
    for ws in wb.worksheets:
        try:
            cabecalho, colunas = localizar_cabecalho(ws)
            session["aba"] = ws.title
            return Path(caminho), wb, ws, cabecalho, colunas
        except ValueError:
            pass
    wb.close()
    raise ValueError("Não encontrei os cabeçalhos esperados em nenhuma aba.")


def elegibilidade(numero, ultima_data, status):
    if not vazio(status):
        return False, "Status preenchido"
    maximo = int(session.get("max_contatos", 2))
    if numero >= maximo:
        return False, "Máximo atingido"
    if numero == 0:
        return True, "Primeiro contato"
    anterior = converter_data(ultima_data)
    if not anterior:
        return False, "Última data não informada"
    proxima = anterior + timedelta(days=int(session.get("dias_segundo", 3)))
    return (True, f"Contato {numero + 1} liberado") if datetime.now() >= proxima else (False, f"Aguardar até {proxima:%d/%m/%Y}")


def listar_registros():
    _, wb, ws, cabecalho, colunas = abrir_planilha(False)
    _, wb_valores, ws_valores, _, _ = abrir_planilha(True)
    registros = []
    try:
        for linha in range(cabecalho + 1, ws.max_row + 1):
            matricula = ws.cell(linha, colunas["matricula"]).value
            msg1 = ws_valores.cell(linha, colunas["mensagem"]).value
            if vazio(matricula) or vazio(msg1):
                continue
            numero = numero_inteiro(ws.cell(linha, colunas["numero"]).value)
            msg2 = ws_valores.cell(linha, colunas["mensagem2"]).value if "mensagem2" in colunas else None
            status = ws.cell(linha, colunas["status"]).value
            ultima = ws.cell(linha, colunas["ultima_data"]).value
            apto, motivo = elegibilidade(numero, ultima, status)
            registros.append({
                "linha": linha,
                "matricula": str(matricula).strip(),
                "nome": ws.cell(linha, colunas["nome"]).value or "",
                "mensagem": msg1 if numero == 0 else (msg2 or msg1),
                "numero": numero,
                "status": status or "",
                "ultima": ultima.strftime("%d/%m/%Y %H:%M") if isinstance(ultima, datetime) else (ultima or ""),
                "apto": apto,
                "motivo": motivo,
            })
    finally:
        wb_valores.close()
        wb.close()
    return registros


def sid():
    if "sid" not in session:
        session["sid"] = secrets.token_urlsafe(24)
    return session["sid"]


def token_cache():
    return TOKEN_CACHES.setdefault(sid(), msal.SerializableTokenCache())


def msal_app(cache=None):
    return msal.ConfidentialClientApplication(
        CLIENT_ID,
        authority=f"https://login.microsoftonline.com/{TENANT_ID}",
        client_credential=CLIENT_SECRET,
        token_cache=cache,
    )


def access_token():
    if not configurado():
        return None
    cache = token_cache()
    cliente = msal_app(cache)
    contas = cliente.get_accounts()
    resultado = cliente.acquire_token_silent(SCOPES, account=contas[0]) if contas else None
    return resultado.get("access_token") if resultado else None


def graph(method, path, token, **kwargs):
    resposta = requests.request(method, GRAPH + path, headers={"Authorization": f"Bearer {token}"}, timeout=30, **kwargs)
    if not resposta.ok:
        try:
            detalhe = resposta.json().get("error", {}).get("message", resposta.text)
        except ValueError:
            detalhe = resposta.text
        raise RuntimeError(f"Microsoft Graph {resposta.status_code}: {detalhe}")
    return resposta.json() if resposta.content else {}


def share_id(url: str) -> str:
    codificado = base64.urlsafe_b64encode(url.encode("utf-8")).decode("ascii").rstrip("=")
    return "u!" + codificado


def conectar_planilha_onedrive(url: str, token: str):
    compartilhamento = share_id(url)
    item = graph("GET", f"/shares/{compartilhamento}/driveItem", token, params={
        "$select": "id,name,webUrl,parentReference,file"
    })
    nome = item.get("name", "planilha.xlsx")
    extensao = Path(nome).suffix.lower()
    if extensao not in {".xlsx", ".xlsm"}:
        raise ValueError("O link compartilhado não aponta para uma planilha .xlsx ou .xlsm.")
    resposta = requests.get(
        GRAPH + f"/shares/{compartilhamento}/driveItem/content",
        headers={"Authorization": f"Bearer {token}"}, timeout=60, allow_redirects=True,
    )
    resposta.raise_for_status()
    destino = DATA_DIR / f"onedrive_{sid()}{extensao}"
    destino.write_bytes(resposta.content)
    session.update(
        planilha=str(destino), nome_original=nome, origem="onedrive",
        onedrive_url=url, onedrive_share_id=compartilhamento,
        onedrive_drive_id=item.get("parentReference", {}).get("driveId"),
        onedrive_item_id=item.get("id"), onedrive_web_url=item.get("webUrl", url),
    )


def sincronizar_onedrive(token: str):
    if session.get("origem") != "onedrive":
        return
    resposta = requests.get(
        GRAPH + f"/shares/{session['onedrive_share_id']}/driveItem/content",
        headers={"Authorization": f"Bearer {token}"}, timeout=60, allow_redirects=True,
    )
    resposta.raise_for_status()
    Path(session["planilha"]).write_bytes(resposta.content)


def salvar_no_onedrive(caminho: Path, token: str):
    if session.get("origem") != "onedrive":
        return
    drive_id = session.get("onedrive_drive_id")
    item_id = session.get("onedrive_item_id")
    if not drive_id or not item_id:
        raise RuntimeError("Identificação do arquivo do OneDrive não encontrada.")
    resposta = requests.put(
        GRAPH + f"/drives/{drive_id}/items/{item_id}/content",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/octet-stream"},
        data=caminho.read_bytes(), timeout=120,
    )
    if not resposta.ok:
        raise RuntimeError(f"Não foi possível atualizar o arquivo no OneDrive: {resposta.text[:500]}")


def localizar_usuario(matricula, token):
    dados = graph("GET", "/users", token, params={
        "$filter": f"employeeId eq '{matricula.replace(chr(39), chr(39) * 2)}'",
        "$select": "id,displayName,userPrincipalName,mail,employeeId",
    })
    usuarios = dados.get("value", [])
    if not usuarios:
        raise RuntimeError(f"Nenhum usuário encontrado com employeeId {matricula}.")
    if len(usuarios) > 1:
        raise RuntimeError(f"Mais de um usuário encontrado com employeeId {matricula}.")
    return usuarios[0]


def enviar_teams(registro, token):
    remetente = graph("GET", "/me", token, params={"$select": "id,displayName,userPrincipalName"})
    destino = localizar_usuario(registro["matricula"], token)
    chat = graph("POST", "/chats", token, json={
        "chatType": "oneOnOne",
        "members": [
            {"@odata.type": "#microsoft.graph.aadUserConversationMember", "roles": ["owner"], "user@odata.bind": f"https://graph.microsoft.com/v1.0/users('{remetente['id']}')"},
            {"@odata.type": "#microsoft.graph.aadUserConversationMember", "roles": ["owner"], "user@odata.bind": f"https://graph.microsoft.com/v1.0/users('{destino['id']}')"},
        ],
    })
    conteudo = html.escape(str(registro["mensagem"])).replace("\n", "<br>")
    mensagem = graph("POST", f"/chats/{chat['id']}/messages", token, json={"body": {"contentType": "html", "content": conteudo}})
    return remetente, destino, chat, mensagem


def confirmar_planilha(linha, remetente, destino, chat, mensagem, token):
    caminho, wb, ws, _, colunas = abrir_planilha(False)
    try:
        atual = numero_inteiro(ws.cell(linha, colunas["numero"]).value)
        backup = caminho.with_name(caminho.stem + "_backup" + caminho.suffix)
        if not backup.exists():
            shutil.copy2(caminho, backup)
        ws.cell(linha, colunas["numero"]).value = atual + 1
        data_cell = ws.cell(linha, colunas["ultima_data"])
        data_cell.value = datetime.now()
        data_cell.number_format = "dd/mm/yyyy hh:mm"
        wb.save(caminho)
        salvar_no_onedrive(caminho, token)
    finally:
        wb.close()


TRANSLATIONS = {
 "br":{"title":"Controle automático pelo Teams","subtitle":"Envio seguro pelo Microsoft Graph em nome do usuário conectado.","entra":"Microsoft Entra","waiting":"Aguardando dados da TI.","fill":"Preencha o arquivo .env e reinicie o aplicativo.","connected":"Conectado como","logout":"Sair","config":"Configuração encontrada. Entre com a conta corporativa.","login":"Entrar com Microsoft","sheet":"Planilha","file":"Arquivo Excel","tab":"Aba opcional","load":"Carregar","rules":"Regras","days":"Dias para novo contato","maximum":"Máximo de contatos","save":"Salvar","contacts":"Contatos","eligible":"apto(s)","send_all":"Enviar todos os aptos","download":"Baixar Excel atualizado","row":"Linha","employee":"Funcionário","message":"Mensagem","control":"Controle","action":"Ação","contact":"Contato","send":"Enviar automaticamente","confirm":"Enviar agora {n} mensagem(ns) pelo Teams? A planilha será atualizada somente para os envios confirmados.","sending":"Enviando... não feche esta página"},
 "esp":{"title":"Control automático por Teams","subtitle":"Envío seguro mediante Microsoft Graph en nombre del usuario conectado.","entra":"Microsoft Entra","waiting":"Esperando datos de TI.","fill":"Complete el archivo .env y reinicie la aplicación.","connected":"Conectado como","logout":"Salir","config":"Configuración encontrada. Ingrese con la cuenta corporativa.","login":"Ingresar con Microsoft","sheet":"Hoja de cálculo","file":"Archivo Excel","tab":"Pestaña opcional","load":"Cargar","rules":"Reglas","days":"Días para nuevo contacto","maximum":"Máximo de contactos","save":"Guardar","contacts":"Contactos","eligible":"habilitado(s)","send_all":"Enviar todos los habilitados","download":"Descargar Excel actualizado","row":"Fila","employee":"Empleado","message":"Mensaje","control":"Control","action":"Acción","contact":"Contacto","send":"Enviar automáticamente","confirm":"¿Enviar ahora {n} mensaje(s) por Teams? La hoja solo se actualizará para los envíos confirmados.","sending":"Enviando... no cierre esta página"},
 "usa":{"title":"Automatic Teams Contact Control","subtitle":"Secure delivery through Microsoft Graph on behalf of the signed-in user.","entra":"Microsoft Entra","waiting":"Waiting for IT credentials.","fill":"Complete the .env file and restart the application.","connected":"Signed in as","logout":"Sign out","config":"Configuration found. Sign in with your corporate account.","login":"Sign in with Microsoft","sheet":"Spreadsheet","file":"Excel file","tab":"Optional worksheet","load":"Load","rules":"Rules","days":"Days before next contact","maximum":"Maximum contacts","save":"Save","contacts":"Contacts","eligible":"eligible","send_all":"Send all eligible","download":"Download updated Excel","row":"Row","employee":"Employee","message":"Message","control":"Control","action":"Action","contact":"Contact","send":"Send automatically","confirm":"Send {n} Teams message(s) now? The spreadsheet will only be updated for confirmed deliveries.","sending":"Sending... do not close this page"}
}

REASON_TRANSLATIONS = {
 "esp":{"Status preenchido":"Estado completado","Máximo atingido":"Máximo alcanzado","Primeiro contato":"Primer contacto","Última data não informada":"Última fecha no informada"},
 "usa":{"Status preenchido":"Status completed","Máximo atingido":"Maximum reached","Primeiro contato":"First contact","Última data não informada":"Last contact date missing"},
}


def traduzir_motivo(motivo, lang):
    if lang == "br":
        return motivo
    tabela = REASON_TRANSLATIONS.get(lang, {})
    if motivo in tabela:
        return tabela[motivo]
    if motivo.startswith("Contato ") and motivo.endswith(" liberado"):
        numero = motivo.split()[1]
        return (f"Contacto {numero} habilitado" if lang == "esp" else f"Contact {numero} eligible")
    if motivo.startswith("Aguardar até "):
        data = motivo.replace("Aguardar até ", "")
        return (f"Esperar hasta {data}" if lang == "esp" else f"Wait until {data}")
    return motivo

HTML = r"""
<!doctype html><html lang="{{'pt-BR' if lang=='br' else ('es' if lang=='esp' else 'en-US')}}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{{txt.title}}</title><style>
:root{--navy:#254a68;--blue:#4d8db4;--sky:#82b8d3;--steel:#8d9aa4;--paper:#eef2f4;--ink:#263843;--green:#4f8a65}*{box-sizing:border-box}body{margin:0;background:linear-gradient(135deg,#eef2f4,#d8e1e6);font:15px "Segoe UI",system-ui;color:var(--ink);min-height:100vh}body:before,body:after{content:"";position:fixed;border-radius:50%;z-index:-1;opacity:.13;background:repeating-conic-gradient(#315b75 0 9deg,transparent 9deg 18deg);-webkit-mask:radial-gradient(circle,#0000 0 34%,#000 35% 55%,#0000 56%)}body:before{width:620px;height:620px;right:-170px;bottom:-190px}body:after{width:330px;height:330px;right:290px;bottom:-100px}.topbar{background:linear-gradient(90deg,#274f6b,#5fa6cb,#2c6688);color:white;box-shadow:0 3px 12px #18384c55}.head{max-width:1200px;margin:auto;padding:21px 24px;display:flex;justify-content:space-between;align-items:center;gap:18px}.brand h1{margin:0;font-size:27px}.brand p{margin:5px 0 0;opacity:.92}.langs{display:flex;gap:6px}.langs a{color:white;text-decoration:none;border:1px solid #ffffff80;border-radius:18px;padding:7px 11px;font-weight:700}.langs a.active{background:white;color:#315e7c}.wrap{max-width:1200px;margin:22px auto;padding:0 18px}.card{background:#f7f9fae8;border:1px solid #b8c4cb;border-radius:16px;padding:19px;margin-bottom:16px;box-shadow:0 7px 20px #29495d18;backdrop-filter:blur(5px)}h2{margin:0 0 17px;color:#315e7c}.grid{display:grid;grid-template-columns:repeat(4,1fr);gap:13px}.wide{grid-column:span 2}label{display:block;font-size:12px;font-weight:700;color:#50636f;margin-bottom:5px}input{width:100%;padding:10px;border:1px solid #aebdc6;border-radius:9px;background:#fff}button,.btn{border:0;border-radius:9px;padding:10px 15px;background:linear-gradient(#579bc2,#36759a);color:white;text-decoration:none;font-weight:700;cursor:pointer;box-shadow:0 2px 5px #254a6833}.success{background:linear-gradient(#64a17a,#477e5b)}.secondary{background:#dce8ee;color:#315e7c;box-shadow:none}.flash{padding:12px 15px;background:#fff3d8;border-left:5px solid #c89232;border-radius:9px;margin-bottom:14px}.warn{border-left:5px solid #c89232}.ok{border-left:5px solid #559271}.muted{color:#667984}.contact-scroll{max-height:430px;overflow:auto;margin-top:12px;border:1px solid #c7d2d9;border-radius:10px;background:#f8fafb}.contact-scroll thead th{position:sticky;top:0;z-index:2}table{width:100%;border-collapse:collapse;table-layout:fixed}th,td{text-align:left;padding:10px;border-bottom:1px solid #cdd6dc;vertical-align:top}th{background:#dce7ec;color:#365c72;font-size:12px}.message-cell .message-preview{display:-webkit-box;-webkit-box-orient:vertical;-webkit-line-clamp:2;overflow:hidden;line-height:1.38;max-height:2.76em}.actions{display:flex;gap:8px;flex-wrap:wrap}button:disabled{opacity:.55;cursor:not-allowed}@media(max-width:800px){.head{align-items:flex-start;flex-direction:column}.grid{grid-template-columns:1fr}.wide{grid-column:auto}.contact-scroll{max-height:360px}.scroll{overflow:auto}table{min-width:850px}}
</style></head><body><header class="topbar"><div class="head"><div class="brand"><h1>{{txt.title}}</h1><p>{{txt.subtitle}}</p></div><nav class="langs"><a class="{{'active' if lang=='br'}}" href="{{url_for('idioma',codigo='br')}}">BR</a><a class="{{'active' if lang=='esp'}}" href="{{url_for('idioma',codigo='esp')}}">ESP</a><a class="{{'active' if lang=='usa'}}" href="{{url_for('idioma',codigo='usa')}}">USA</a></nav></div></header><main class="wrap">
{% for m in get_flashed_messages() %}<div class="flash">{{m}}</div>{% endfor %}<section class="card {{'ok' if config_ok else 'warn'}}"><h2>1. {{txt.entra}}</h2>{% if not config_ok %}<p><b>{{txt.waiting}}</b> {{txt.fill}}</p><div class="grid"><div><label>Tenant ID</label><input disabled placeholder="Pendente"></div><div><label>Client ID</label><input disabled placeholder="Pendente"></div><div><label>Client Secret</label><input disabled placeholder="Pendente"></div><div><label>Redirect URI</label><input disabled value="{{redirect_uri}}"></div></div>{% elif usuario %}<p>{{txt.connected}} <b>{{usuario}}</b></p><a class="btn secondary" href="{{url_for('logout')}}">{{txt.logout}}</a>{% else %}<p>{{txt.config}}</p><a class="btn" href="{{url_for('login')}}">{{txt.login}}</a>{% endif %}</section>
<section class="card"><h2>2. {{txt.sheet}} — OneDrive / SharePoint / Teams</h2><form action="{{url_for('conectar_compartilhada')}}" method="post" class="grid"><div class="wide"><label>{{'Link compartilhado do OneDrive, SharePoint ou Teams' if lang=='br' else ('Enlace compartido de OneDrive, SharePoint o Teams' if lang=='esp' else 'Shared OneDrive, SharePoint, or Teams link')}}</label><input type="url" name="url_compartilhada" value="{{session.get('onedrive_url',onedrive_default)}}" placeholder="https://..." required></div><div><label>{{txt.tab}}</label><input name="aba_compartilhada" value="{{session.get('aba','')}}" placeholder="COCKPIT SPLUNK"></div><div style="align-self:end"><button>{{'Conectar planilha' if lang=='br' else ('Conectar hoja' if lang=='esp' else 'Connect spreadsheet')}}</button></div></form>{% if session.get('origem')=='onedrive' %}<p class="muted"><b>{{'Arquivo compartilhado ativo' if lang=='br' else ('Archivo compartido activo' if lang=='esp' else 'Active shared file')}}:</b> {{session.get('nome_original')}} · {{session.get('aba')}}</p>{% endif %}</section>
<section class="card"><h2>2B. {{'Arquivo local (somente teste)' if lang=='br' else ('Archivo local (solo prueba)' if lang=='esp' else 'Local file (testing only)')}}</h2><form action="{{url_for('carregar')}}" method="post" enctype="multipart/form-data" class="grid"><div class="wide"><label>{{txt.file}}</label><input type="file" name="arquivo" accept=".xlsx,.xlsm" required></div><div><label>{{txt.tab}}</label><input name="aba" placeholder="COCKPIT SPLUNK"></div><div style="align-self:end"><button>{{txt.load}}</button></div></form>{% if session.get('origem')=='local' %}<p class="muted">{{session.get('nome_original')}} · {{session.get('aba')}}</p>{% endif %}</section>
<section class="card"><h2>3. {{txt.rules}}</h2><form action="{{url_for('configurar')}}" method="post" class="grid"><div><label>{{txt.days}}</label><input type="number" min="0" name="dias_segundo" value="{{session.get('dias_segundo',3)}}"></div><div><label>{{txt.maximum}}</label><input type="number" min="1" name="max_contatos" value="{{session.get('max_contatos',2)}}"></div><div style="align-self:end"><button>{{txt.save}}</button></div></form></section>
{% if registros is not none %}<section class="card"><div style="display:flex;justify-content:space-between;gap:12px;align-items:center"><div><h2>4. {{txt.contacts}}</h2><span class="muted">{{aptos}} {{txt.eligible}}</span></div><div class="actions">{% if aptos > 0 %}<form action="{{url_for('enviar_todos')}}" method="post" onsubmit="return confirmarLote(this, {{aptos}})"><button class="success" {{'disabled' if not usuario else ''}}>{{txt.send_all}} ({{aptos}})</button></form>{% endif %}<a class="btn secondary" href="{{url_for('baixar')}}">{{txt.download}}</a></div></div><div class="contact-scroll"><table><colgroup><col style="width:6%"><col style="width:19%"><col style="width:37%"><col style="width:20%"><col style="width:18%"></colgroup><thead><tr><th>{{txt.row}}</th><th>{{txt.employee}}</th><th>{{txt.message}}</th><th>{{txt.control}}</th><th>{{txt.action}}</th></tr></thead><tbody>{% for r in registros %}<tr><td>{{r.linha}}</td><td><b>{{r.matricula}}</b><br>{{r.nome}}</td><td class="message-cell" title="{{r.mensagem}}"><div class="message-preview">{{r.mensagem}}</div></td><td>{{r.motivo}}<br><span class="muted">{{txt.contact}} {{r.numero}} · {{r.ultima}}</span></td><td>{% if r.apto %}<form action="{{url_for('enviar',linha=r.linha)}}" method="post" onsubmit="this.querySelector('button').disabled=true"><button class="success" {{'disabled' if not usuario else ''}}>{{txt.send}}</button></form>{% endif %}</td></tr>{% endfor %}</tbody></table></div></section>{% endif %}</main><script>const confirmTemplate={{txt.confirm|tojson}},sendingText={{txt.sending|tojson}};function confirmarLote(f,n){const ok=confirm(confirmTemplate.replace('{n}',n));if(ok){const b=f.querySelector('button');b.disabled=true;b.textContent=sendingText}return ok}</script></body></html>
"""


@app.get("/")
def inicio():
    lang = session.get("lang", "br")
    txt = dict(TRANSLATIONS.get(lang, TRANSLATIONS["br"]))
    if session.get("origem") == "onedrive":
        txt["download"] = {"br": "Abrir planilha compartilhada", "esp": "Abrir hoja compartida", "usa": "Open shared spreadsheet"}.get(lang, "Abrir planilha compartilhada")
    token = access_token()
    usuario = None
    if token:
        try:
            usuario = graph("GET", "/me", token, params={"$select": "displayName,userPrincipalName"}).get("displayName")
            if session.get("origem") == "onedrive":
                sincronizar_onedrive(token)
        except Exception:
            pass
    registros = None
    if session.get("planilha"):
        try:
            registros = listar_registros()
            for registro in registros:
                registro["motivo"] = traduzir_motivo(registro["motivo"], lang)
        except Exception as exc:
            flash(f"Erro ao ler a planilha: {exc}")
    return render_template_string(HTML, config_ok=configurado(), redirect_uri=REDIRECT_URI, onedrive_default=ONEDRIVE_SHARE_URL, usuario=usuario, registros=registros, aptos=sum(r["apto"] for r in registros or []), lang=lang, txt=txt)


@app.get("/idioma/<codigo>")
def idioma(codigo):
    session["lang"] = codigo if codigo in TRANSLATIONS else "br"
    return redirect(url_for("inicio"))


@app.get("/login")
def login():
    if not configurado():
        flash("Preencha os dados da TI no arquivo .env.")
        return redirect(url_for("inicio"))
    fluxo = msal_app(token_cache()).initiate_auth_code_flow(SCOPES, redirect_uri=REDIRECT_URI)
    session["auth_flow"] = fluxo
    return redirect(fluxo["auth_uri"])


@app.get("/auth/callback")
def callback():
    try:
        resultado = msal_app(token_cache()).acquire_token_by_auth_code_flow(session.get("auth_flow", {}), request.args)
        if "access_token" not in resultado:
            raise RuntimeError(resultado.get("error_description", "Login não concluído."))
        flash("Login Microsoft concluído.")
    except Exception as exc:
        flash(f"Falha no login: {exc}")
    return redirect(url_for("inicio"))


@app.get("/logout")
def logout():
    TOKEN_CACHES.pop(session.get("sid", ""), None)
    session.pop("auth_flow", None)
    flash("Sessão Microsoft encerrada.")
    return redirect(url_for("inicio"))


@app.post("/carregar")
def carregar():
    arquivo = request.files.get("arquivo")
    if not arquivo or Path(arquivo.filename).suffix.lower() not in {".xlsx", ".xlsm"}:
        flash("Selecione um arquivo .xlsx ou .xlsm.")
        return redirect(url_for("inicio"))
    ext = Path(arquivo.filename).suffix.lower()
    destino = DATA_DIR / f"trabalho_{secrets.token_hex(5)}{ext}"
    arquivo.save(destino)
    session.update(planilha=str(destino), nome_original=secure_filename(arquivo.filename), aba=request.form.get("aba", "").strip(), origem="local")
    try:
        _, wb, ws, _, _ = abrir_planilha(False)
        session["aba"] = ws.title
        wb.close()
        flash(f"Planilha carregada. Aba: {ws.title}.")
    except Exception as exc:
        session.pop("planilha", None)
        destino.unlink(missing_ok=True)
        flash(f"Não foi possível carregar: {exc}")
    return redirect(url_for("inicio"))


@app.post("/conectar-compartilhada")
def conectar_compartilhada():
    token = access_token()
    if not token:
        flash("Entre com a conta Microsoft antes de conectar a planilha compartilhada.")
        return redirect(url_for("login"))
    url = request.form.get("url_compartilhada", "").strip()
    if not url.startswith("https://"):
        flash("Informe um link HTTPS compartilhado do OneDrive ou SharePoint.")
        return redirect(url_for("inicio"))
    try:
        conectar_planilha_onedrive(url, token)
        session["aba"] = request.form.get("aba_compartilhada", "").strip()
        _, wb, ws, _, _ = abrir_planilha(False)
        session["aba"] = ws.title
        wb.close()
        flash(f"Planilha compartilhada conectada. Aba: {ws.title}.")
    except Exception as exc:
        flash(f"Não foi possível conectar a planilha compartilhada: {exc}")
    return redirect(url_for("inicio"))


@app.post("/configurar")
def configurar():
    session["dias_segundo"] = max(0, int(request.form.get("dias_segundo", 3)))
    session["max_contatos"] = max(1, int(request.form.get("max_contatos", 2)))
    flash("Regras salvas.")
    return redirect(url_for("inicio"))


@app.post("/enviar/<int:linha>")
def enviar(linha):
    token = access_token()
    if not token:
        flash("Faça login com a conta Microsoft.")
        return redirect(url_for("login"))
    sincronizar_onedrive(token)
    registro = next((r for r in listar_registros() if r["linha"] == linha), None)
    if not registro or not registro["apto"]:
        flash("Contato não está liberado.")
        return redirect(url_for("inicio"))
    try:
        remetente, destino, chat, mensagem = enviar_teams(registro, token)
        confirmar_planilha(linha, remetente, destino, chat, mensagem, token)
        flash(f"Mensagem enviada para {destino['displayName']} ({destino['employeeId']}). Planilha atualizada.")
    except Exception as exc:
        flash(f"Envio não realizado; a planilha não foi alterada. Erro: {exc}")
    return redirect(url_for("inicio"))


@app.post("/enviar-todos")
def enviar_todos():
    token = access_token()
    if not token:
        flash("Faça login com a conta Microsoft.")
        return redirect(url_for("login"))

    sincronizar_onedrive(token)

    pendentes = [registro for registro in listar_registros() if registro["apto"]]
    if not pendentes:
        flash("Não existem contatos aptos para envio.")
        return redirect(url_for("inicio"))

    enviados = []
    erros = []
    for registro in pendentes:
        try:
            remetente, destino, chat, mensagem = enviar_teams(registro, token)
            confirmar_planilha(registro["linha"], remetente, destino, chat, mensagem, token)
            enviados.append(f"{registro['matricula']} - {destino['displayName']}")
        except Exception as exc:
            erros.append(f"{registro['matricula']}: {exc}")

    resumo = f"Processamento concluído: {len(enviados)} enviado(s) e {len(erros)} erro(s)."
    if erros:
        resumo += " Falhas: " + " | ".join(erros[:5])
        if len(erros) > 5:
            resumo += f" | e mais {len(erros) - 5} erro(s)."
    flash(resumo)
    return redirect(url_for("inicio"))


@app.get("/baixar")
def baixar():
    if session.get("origem") == "onedrive":
        return redirect(session.get("onedrive_web_url") or session.get("onedrive_url"))
    caminho = Path(session["planilha"])
    original = Path(session.get("nome_original", "planilha.xlsx"))
    return send_file(caminho, as_attachment=True, download_name=f"{original.stem}_atualizada{original.suffix}")


if __name__ == "__main__":
    webbrowser.open_new(f"http://localhost:{PORT}")
    app.run(host="127.0.0.1", port=PORT, debug=False)
